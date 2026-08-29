"""Shared Excel styling & layout for the bot's reports.

Encapsulates the formatting conventions the owner settled on (see the
report-formatting-conventions memory): modern Office "Background 2" tint colors,
zebra rows, muted zeros, an offset so tables start at B2 (empty column A / row 1
for spacing), thin borders, currency/hours number formats, and page setup with a
$B$2-anchored print area + repeated header rows. Report modules (hipp_report, and
the upcoming payroll reports) build their layouts on these primitives.

Colors are explicit hex (not theme references) so they render identically
regardless of the embedded theme; `apply_modern_theme` additionally embeds the
modern Office theme so the in-Excel color-picker palette matches the owner's other
spreadsheets (openpyxl bundles the older theme).
"""

from __future__ import annotations

import re

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

# Offset so the top-left of every table sits at B2 (empty col A + row 1 for spacing).
COL_SHIFT = 1
ROW_SHIFT = 1

HDR = Font(name="Arial", size=10, bold=True)
BODY = Font(name="Arial", size=10)
TITLE = Font(name="Arial", size=14, bold=True)
ZERO = Font(name="Arial", size=10, color="C5C3C3")        # muted zeros
CUR = '$#,##0.00'
HRS = '0.00'
HRS_TOTAL = '0.00 "hrs";[Red]-0.00 "hrs";"-"'              # hours col: dash for zero
_thin = Side(style="thin")
HFILL = PatternFill("solid", fgColor="D9E1F2")            # header
WHITE = PatternFill("solid", fgColor="FFFFFF")            # zebra: white row
ALT = PatternFill("solid", fgColor="DCDADA")             # White, Background 2, Darker 5%
TOTAL_FILL = PatternFill("solid", fgColor="C5C3C3")       # White, Background 2, Darker 15%
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
MID = Alignment(horizontal="center", vertical="center")   # centered data (no wrap)
RIGHT = Alignment(horizontal="right")


def shift_col(letter: str) -> str:
    return get_column_letter(column_index_from_string(letter) + COL_SHIFT)


def addr(col: str, row: int) -> str:
    return f"{shift_col(col)}{row + ROW_SHIFT}"


def border(sides) -> Border:
    """Border with only the named sides ('left'/'right'/'top'/'bottom') drawn thin."""
    return Border(**{s: _thin for s in sides})


def cell(ws, col, row, value, font=BODY, fmt="General", align=None):
    """Plain cell (letterhead) — no fill, no border. Logical (col, row); offset applied."""
    c = ws[addr(col, row)]
    c.value = value
    c.font = font
    c.number_format = fmt
    if align:
        c.alignment = align
    return c


def put(ws, col, r, value, fmt, align, fill, sides, base_font=BODY, zero_gray=False):
    """Write a table cell with fill + border; grays a zero-valued numeric cell."""
    c = ws[addr(col, r)]
    c.value = value
    c.font = ZERO if (zero_gray and isinstance(value, (int, float)) and value == 0) else base_font
    c.number_format = fmt
    c.fill = fill
    c.border = border(sides)
    if align:
        c.alignment = align
    return c


def finalize(ws, header_last_row, last_col, last_row, landscape):
    """Print area = $B$2 -> bottom-right edited cell; repeat the header rows; fit to
    width; modest margins. `last_col` is the logical rightmost column, `last_row` the
    actual bottom row, `header_last_row` the logical header row."""
    ws.print_area = f"${shift_col('A')}${1 + ROW_SHIFT}:${shift_col(last_col)}${last_row}"
    ws.print_title_rows = f"{1 + ROW_SHIFT}:{header_last_row + ROW_SHIFT}"
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.5, bottom=0.5, header=0.3, footer=0.3)
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def _modern_theme() -> bytes | None:
    """openpyxl bundles the OLD Office 2007-2010 theme, so downloaded files show that
    palette (beige "Background 2", muted accents) instead of the modern one the owner
    uses. Rebuild the theme with the modern Office color scheme (+ Calibri Light
    headings). Returns theme bytes, or None if unavailable."""
    try:
        from openpyxl.writer.theme import theme_xml
    except Exception:  # noqa: BLE001 - never let theming break report generation
        return None
    xml = theme_xml.decode() if isinstance(theme_xml, (bytes, bytearray)) else theme_xml
    modern_clr = (
        '<a:clrScheme name="Office">'
        '<a:dk1><a:sysClr val="windowText" lastClr="000000"/></a:dk1>'
        '<a:lt1><a:sysClr val="window" lastClr="FFFFFF"/></a:lt1>'
        '<a:dk2><a:srgbClr val="44546A"/></a:dk2>'
        '<a:lt2><a:srgbClr val="E7E6E6"/></a:lt2>'
        '<a:accent1><a:srgbClr val="4472C4"/></a:accent1>'
        '<a:accent2><a:srgbClr val="ED7D31"/></a:accent2>'
        '<a:accent3><a:srgbClr val="A5A5A5"/></a:accent3>'
        '<a:accent4><a:srgbClr val="FFC000"/></a:accent4>'
        '<a:accent5><a:srgbClr val="5B9BD5"/></a:accent5>'
        '<a:accent6><a:srgbClr val="70AD47"/></a:accent6>'
        '<a:hlink><a:srgbClr val="0563C1"/></a:hlink>'
        '<a:folHlink><a:srgbClr val="954F72"/></a:folHlink>'
        '</a:clrScheme>'
    )
    xml = re.sub(r"<a:clrScheme.*?</a:clrScheme>", lambda _m: modern_clr, xml, count=1, flags=re.S)
    xml = xml.replace('<a:latin typeface="Cambria"/>', '<a:latin typeface="Calibri Light"/>', 1)
    return xml.encode("utf-8")


_MODERN_THEME = _modern_theme()


def apply_modern_theme(wb):
    """Embed the modern Office theme so the in-Excel palette matches the owner's sheets."""
    if _MODERN_THEME:
        wb.loaded_theme = _MODERN_THEME
